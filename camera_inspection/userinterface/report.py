from PySide6.QtWidgets import (
                            QWidget,
                            QVBoxLayout,
                            QHBoxLayout,
                            QDateEdit,
                            QComboBox,
                            QLabel,
                            QLineEdit,
                            QPushButton,
                            QTableWidget,
                            QHeaderView,
                            QTableWidgetItem,
                            QFileDialog,
                            QDialog
                            )
from PySide6.QtCore import QDate,QTimer,Qt
from PySide6.QtGui import QColor,QPixmap
from datetime import datetime,time

from openpyxl.styles import PatternFill
from openpyxl import Workbook


from db.db_queries import fetch_report


class Report(QWidget):
    def __init__(self):
        super().__init__()
        main = QVBoxLayout(self)

        # ---------- HEADER ----------
        header = QHBoxLayout()
        left = QHBoxLayout()

        self.from_dt = QDateEdit(calendarPopup=True)
        self.from_dt.setDate(QDate.currentDate().addDays(-7))

        self.to_dt = QDateEdit(calendarPopup=True)
        self.to_dt.setDate(QDate.currentDate())

        self.status = QComboBox()
        self.status.addItems(["ALL", "OK", "NOT_OK"])

        #  debounce timer
        self.search_timer = QTimer(self)
        self.search_timer.setSingleShot(True)
        self.search_timer.timeout.connect(self.reset_and_load)

        left.addWidget(QLabel("From"))
        left.addWidget(self.from_dt)
        left.addWidget(QLabel("To"))
        left.addWidget(self.to_dt)
        left.addWidget(QLabel("Status"))
        left.addWidget(self.status)
        left.addWidget(QLabel("Unique No"))
        self.search_unique = QLineEdit()
        left.addWidget(self.search_unique)

        right = QHBoxLayout()
        self.btn_excel = QPushButton("Export Excel")
        self.btn_excel.setStyleSheet(
            "background:#28a745;color:white;font-weight:bold;padding:6px 18px;"
        )
        right.addStretch()
        right.addWidget(self.btn_excel)

        header.addLayout(left)
        header.addStretch()
        header.addLayout(right)
        main.addLayout(header)

        # ---------- PAGINATION ----------
        pager = QHBoxLayout()
        self.btn_prev = QPushButton("⬅ Previous")
        self.btn_next = QPushButton("Next ➡")
        self.page_lbl = QLabel("Page 1")

        pager.addStretch()
        pager.addWidget(self.btn_prev)
        pager.addWidget(self.page_lbl)
        pager.addWidget(self.btn_next)
        pager.addStretch()

        main.addLayout(pager)

        self.page_size = 10
        self.current_page = 1
        self.total_rows = 0

        # ---------- TABLE ----------
        self.table = QTableWidget(0, 10)
        self.table.setHorizontalHeaderLabels(
            [
                "Employee ID",
                "Work Order",
                "Charge No",
                "Serial No",
                "Vendor Code",
                "Unique No",
                "Image",
                "Status",
                "Date",
                "Time",
            ]
        )

        self.table.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
        self.table.verticalHeader().setVisible(False)

        from PySide6.QtWidgets import QAbstractItemView

        self.table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self.table.setSelectionBehavior(QAbstractItemView.SelectRows)
        self.table.setSelectionMode(QAbstractItemView.SingleSelection)

        main.addWidget(self.table)

        # ---------- SIGNALS ----------
        self.from_dt.dateChanged.connect(self.on_filter_changed)
        self.to_dt.dateChanged.connect(self.on_filter_changed)
        self.status.currentIndexChanged.connect(self.on_filter_changed)

        self.search_unique.textChanged.connect(self.on_search_text_changed)

        self.btn_prev.clicked.connect(self.prev_page)
        self.btn_next.clicked.connect(self.next_page)
        self.btn_excel.clicked.connect(self.export_excel)

        self.load()

    def on_search_text_changed(self, text):
        # restart debounce timer on each key press
        self.search_timer.start(400)  # milliseconds

    def on_filter_changed(self, *args):
        # debounce all filters
        self.search_timer.start(400)

    # ---------- LOAD ----------
    def load(self):
        self.table.setRowCount(0)

        f = datetime.combine(self.from_dt.date().toPython(), time.min)
        t = datetime.combine(self.to_dt.date().toPython(), time.max)
        unique_text = self.search_unique.text().strip()
        offset = (self.current_page - 1) * self.page_size

        rows, self.total_rows = fetch_report(
            f, t, self.status.currentText(), unique_text, self.page_size, offset
        )

        for r in rows:
            row = self.table.rowCount()
            self.table.insertRow(row)
            self.table.setRowHeight(row, 150)

            for c in range(6):
                item = QTableWidgetItem(str(r[c]))
                item.setTextAlignment(Qt.AlignCenter)
                self.table.setItem(row, c, item)

            status_item = QTableWidgetItem(r[7])
            status_item.setTextAlignment(Qt.AlignCenter)
            status_item.setForeground(
                QColor("green") if r[7] == "OK" else QColor("red")
            )
            self.table.setItem(row, 7, status_item)

            self.table.setItem(row, 8, QTableWidgetItem(r[8].strftime("%Y-%m-%d")))
            self.table.setItem(row, 9, QTableWidgetItem(r[8].strftime("%H:%M:%S")))

            QTimer.singleShot(10, lambda row=row, img=r[6]: self._set_image(row, img))

        total_pages = max(1, (self.total_rows + self.page_size - 1) // self.page_size)
        self.page_lbl.setText(f"Page {self.current_page} / {total_pages}")

        self.btn_prev.setEnabled(self.current_page > 1)
        self.btn_next.setEnabled(self.current_page < total_pages)

    def reset_and_load(self):
        self.current_page = 1
        self.load()

    def next_page(self):
        self.current_page += 1
        self.load()

    def prev_page(self):
        if self.current_page > 1:
            self.current_page -= 1
            self.load()

    def _set_image(self, row, img_bytes):
        pix = QPixmap()
        pix.loadFromData(bytes(img_bytes))

        lbl = QLabel(alignment=Qt.AlignCenter)
        lbl.setPixmap(pix.scaled(220, 130, Qt.KeepAspectRatio))
        lbl.setCursor(Qt.PointingHandCursor)

        # 👇 THIS IS THE KEY LINE
        lbl.mousePressEvent = lambda e, b=img_bytes: self.show_image_dialog(b)

        self.table.setCellWidget(row, 6, lbl)

    def export_excel(self):
        path, _ = QFileDialog.getSaveFileName(
            self, "Save Excel", "", "Excel Files (*.xlsx)"
        )
        if not path:
            return

        if not path.lower().endswith(".xlsx"):
            path += ".xlsx"

        wb = Workbook()
        ws = wb.active
        ws.title = "Inspection Report"

        ws.append(
            [
                "Employee ID",
                "Work Order",
                "Charge No",
                "Serial No",
                "Part No",
                "Unique No",
                "Status",
                "Date",
                "Time",
            ]
        )

        green = PatternFill("solid", fgColor="C6EFCE")
        red = PatternFill("solid", fgColor="FFC7CE")

        f = datetime.combine(self.from_dt.date().toPython(), time.min)
        t = datetime.combine(self.to_dt.date().toPython(), time.max)
        unique_text = self.search_unique.text().strip()

        rows, _ = fetch_report(f, t, self.status.currentText(), unique_text)

        for r in rows:
            ws.append(
                [
                    r[0],
                    r[1],
                    r[2],
                    r[3],
                    r[4],
                    r[5],
                    r[7],
                    r[8].strftime("%Y-%m-%d"),
                    r[8].strftime("%H:%M:%S"),
                ]
            )

            row_idx = ws.max_row
            ws[f"G{row_idx}"].fill = green if r[7] == "OK" else red

        wb.save(path)

    def show_image_dialog(self, front_bytes, back_bytes):
        dlg = QDialog(self)
        dlg.setWindowTitle("Front / Back View")
        dlg.resize(1000, 500)

        front = QPixmap()
        back = QPixmap()
        front.loadFromData(bytes(front_bytes))
        back.loadFromData(bytes(back_bytes))

        lbl_f = QLabel()
        lbl_b = QLabel()

        lbl_f.setPixmap(front.scaled(480, 480, Qt.KeepAspectRatio))
        lbl_b.setPixmap(back.scaled(480, 480, Qt.KeepAspectRatio))

        lay = QHBoxLayout(dlg)
        lay.addWidget(lbl_f)
        lay.addWidget(lbl_b)

        dlg.exec()
